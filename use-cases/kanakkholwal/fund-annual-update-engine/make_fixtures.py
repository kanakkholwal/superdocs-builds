"""Builds last year's fund document set and this year's figures. All fictional.

Meridian Global Equity Fund does not exist and every figure here is invented.
"""

import csv
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
BASELINE = ROOT / "fixtures" / "baseline"
DATA = ROOT / "fixtures" / "data"

FUND = "Meridian Global Equity Fund"
ISIN = "IE00BMERID01"


def _table(doc, name, headers, rows):
    doc.add_paragraph(name, style="Heading 2")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, headers, strict=True):
        cell.text = head
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = str(value)
    doc.add_paragraph()


def build_prospectus(path: Path) -> None:
    doc = Document()
    doc.add_heading(FUND, level=0)
    doc.add_paragraph("Prospectus Supplement dated 1 April 2025")

    doc.add_paragraph("1. The Fund", style="Heading 1")
    doc.add_paragraph(
        "The Fund is a sub-fund of Meridian Investment Funds ICAV, an open-ended umbrella "
        "investment company with variable capital and segregated liability between sub-funds, "
        f"authorised in Ireland. The ISIN of the accumulating share class is {ISIN}."
    )
    doc.add_paragraph(
        "Investors should note that the value of an investment in the Fund and the income derived "
        "from it may go down as well as up, and an investor may not get back the amount originally "
        "invested. There is no guarantee that the investment objective will be achieved."
    )

    doc.add_paragraph("2. Investment objective and policy", style="Heading 1")
    doc.add_paragraph(
        "The Fund seeks long-term capital growth by investing primarily in equity securities of "
        "companies listed in developed markets worldwide. The Fund is actively managed and is not "
        "constrained by any benchmark."
    )
    doc.add_paragraph(
        "The Manager may, at its discretion, hold ancillary liquid assets and may use financial "
        "derivative instruments for efficient portfolio management, hedging and investment "
        "purposes, in each case subject to the conditions and limits laid down by the Central Bank."
    )

    doc.add_paragraph("3. Fees and expenses", style="Heading 1")
    _table(
        doc,
        "Charges",
        ["Charge", "Amount"],
        [
            ["Ongoing charges figure", "0.82%"],
            ["Management fee", "0.65%"],
            ["Entry charge", "Nil"],
            ["Exit charge", "Nil"],
        ],
    )
    doc.add_paragraph(
        "The ongoing charges figure is based on expenses for the twelve month period ending "
        "31 December 2024 and may vary from year to year."
    )
    doc.save(path)


def build_factsheet(path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)

    doc.add_heading(FUND, level=0)
    doc.add_paragraph("Monthly factsheet. Data as at 31 December 2024")

    doc.add_paragraph("Key facts", style="Heading 1")
    _table(
        doc,
        "Overview",
        ["Item", "Value"],
        [
            ["Fund size", "USD 1,284.6m"],
            ["NAV per unit", "USD 18.42"],
            ["Ongoing charges figure", "0.82%"],
            ["ISIN", ISIN],
        ],
    )

    _table(
        doc,
        "Performance",
        ["Calendar year", "Fund return"],
        [
            ["2021", "18.4%"],
            ["2022", "-14.2%"],
            ["2023", "21.7%"],
            ["2024", "11.9%"],
            ["2025", "n/a"],
        ],
    )
    doc.add_paragraph(
        "Past performance is not a reliable indicator of future results. Returns are shown net of "
        "ongoing charges and assume reinvestment of income."
    )

    _table(
        doc,
        "Holdings",
        ["Rank", "Holding", "Weight"],
        [
            ["1", "Aurora Semiconductor", "4.8%"],
            ["2", "Northwind Logistics", "4.1%"],
            ["3", "Calder Pharmaceuticals", "3.7%"],
            ["4", "Basalt Energy", "3.3%"],
            ["5", "Vantage Software", "3.1%"],
        ],
    )
    doc.save(path)


def build_kid(path: Path) -> None:
    doc = Document()
    doc.add_heading("Key Information Document", level=0)
    doc.add_paragraph(
        "You are about to purchase a product that is not simple and may be difficult to understand."
    )
    doc.add_paragraph(f"Product: {FUND}. ISIN: {ISIN}. Data as at 31 December 2024")

    doc.add_paragraph("What is this product?", style="Heading 1")
    doc.add_paragraph(
        "Type: an open-ended UCITS sub-fund. Objective: long-term capital growth from a global "
        "portfolio of developed market equities."
    )

    doc.add_paragraph("What are the risks and what could I get in return?", style="Heading 1")
    doc.add_paragraph("Summary risk indicator: 4 out of 7")
    doc.add_paragraph(
        "This product does not include any protection from future market performance, so you could "
        "lose some or all of your investment."
    )
    _table(
        doc,
        "Scenarios",
        ["Scenario", "1 year", "5 years"],
        [
            ["Stress", "USD 4,120", "USD 5,980"],
            ["Unfavourable", "USD 8,640", "USD 9,510"],
            ["Moderate", "USD 10,540", "USD 13,720"],
            ["Favourable", "USD 13,180", "USD 18,040"],
        ],
    )

    doc.add_paragraph("What are the costs?", style="Heading 1")
    _table(
        doc,
        "Costs",
        ["Cost", "Amount"],
        [
            ["Ongoing charges figure", "0.82%"],
            ["Entry costs", "0.00%"],
            ["Exit costs", "0.00%"],
            ["Transaction costs", "0.09%"],
        ],
    )
    doc.save(path)


# holding_2 deliberately omits a weight: the engine must flag it rather than invent one.
NEW_FIGURES = [
    ("ocf", "0.79%", "2025 expense ledger", "AR-2025 Charges B14"),
    ("sri", "5", "PRIIPs SRI calculation", "RISK-2025 SRI C7"),
    ("fund_size_musd", "USD 1,402.3m", "Dec-2025 valuation", "NAV-2025 AUM D31"),
    ("nav_per_unit", "USD 20.71", "Dec-2025 valuation", "NAV-2025 NAV E31"),
    ("as_of_date", "31 December 2025", "Reporting calendar", "CAL-2025 A2"),
    ("perf_2021", "18.4%", "Performance book", "PERF-2025 CY B2"),
    ("perf_2022", "-14.2%", "Performance book", "PERF-2025 CY B3"),
    ("perf_2023", "21.7%", "Performance book", "PERF-2025 CY B4"),
    ("perf_2024", "11.9%", "Performance book", "PERF-2025 CY B5"),
    ("perf_2025", "9.6%", "Performance book", "PERF-2025 CY B6"),
    ("holding_1", "Aurora Semiconductor|5.2%", "Dec-2025 holdings", "HLD-2025 row 1"),
    ("holding_2", "Vantage Software", "Dec-2025 holdings", "HLD-2025 row 2"),
    ("holding_3", "Northwind Logistics|3.9%", "Dec-2025 holdings", "HLD-2025 row 3"),
    ("holding_4", "Calder Pharmaceuticals|3.5%", "Dec-2025 holdings", "HLD-2025 row 4"),
    ("holding_5", "Meridian Water Utilities|3.0%", "Dec-2025 holdings", "HLD-2025 row 5"),
]


def build_data(path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(["field_id", "new_value", "source_document", "source_cell"])
        writer.writerows(NEW_FIGURES)


def main() -> int:
    BASELINE.mkdir(parents=True, exist_ok=True)
    DATA.mkdir(parents=True, exist_ok=True)
    build_prospectus(BASELINE / "prospectus.docx")
    build_factsheet(BASELINE / "factsheet.docx")
    build_kid(BASELINE / "kid.docx")
    build_data(DATA / "figures-2025.csv")
    for path in sorted(BASELINE.iterdir()) + sorted(DATA.iterdir()):
        print(f"  {path.relative_to(ROOT)}  {path.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
