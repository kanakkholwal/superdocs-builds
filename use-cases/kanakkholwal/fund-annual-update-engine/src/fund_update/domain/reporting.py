"""The two documents a compliance reviewer actually reads: the marked version and the change log."""

import html
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt

from fund_update.domain.guards import GuardResult
from fund_update.domain.plan import Plan

STYLE = """
:root{color-scheme:light dark;
--bg:light-dark(#fbfbfd,#16171c);--card:light-dark(#fff,#1e2026);--line:light-dark(#e3e5ea,#2c2f38);
--ink:light-dark(#1d2027,#e9eaee);--dim:light-dark(#666b78,#a2a7b4);
--add:light-dark(#0a7c42,#54d18a);--del:light-dark(#b3261e,#ff8a80);--warn:light-dark(#8a5a00,#e5b567)}
*{box-sizing:border-box}
body{margin:0;padding:2rem 1.25rem;background:var(--bg);color:var(--ink);
font:15px/1.55 ui-sans-serif,system-ui,-apple-system,"Segoe UI",sans-serif}
.wrap{max-width:64rem;margin:0 auto}
h1{font-size:1.35rem;margin:0 0 .25rem}
h2{font-size:1rem;margin:2rem 0 .75rem;text-transform:uppercase;letter-spacing:.04em;color:var(--dim)}
.sub{color:var(--dim);margin:0 0 1.5rem}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;
  padding:.9rem 1rem;margin:.5rem 0}
.row{display:flex;flex-wrap:wrap;gap:.5rem;align-items:baseline}
.label{font-weight:600}
.doc{font-size:.75rem;color:var(--dim);border:1px solid var(--line);border-radius:99px;padding:.05rem .5rem}
.old{color:var(--del);text-decoration:line-through;font-variant-numeric:tabular-nums}
.new{color:var(--add);font-weight:600;font-variant-numeric:tabular-nums}
.arrow{color:var(--dim)}
.cite{font-size:.8rem;color:var(--dim);margin-top:.35rem}
.cite code{background:transparent;color:inherit}
.ok{color:var(--add)}.bad{color:var(--del)}.warn{color:var(--warn)}
ul{margin:.4rem 0;padding-left:1.1rem}
li{margin:.15rem 0}
.tot{display:flex;gap:1.5rem;flex-wrap:wrap;margin:.5rem 0 1.5rem;color:var(--dim);font-size:.9rem}
.tot b{color:var(--ink);font-size:1.05rem}
@media print{body{padding:0}.card{break-inside:avoid}}
"""


def write_review_html(
    plan: Plan, guards: dict[str, GuardResult], path: Path, fund: str = "Meridian Global Equity Fund"
) -> Path:
    """The marked version. Grouped by document, each change beside the cell it came from, so a
    reviewer can sign it off without opening the data file."""
    parts: list[str] = [
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width,initial-scale=1'>",
        "<meta name='color-scheme' content='light dark'>",
        f"<title>{html.escape(fund)} annual refresh</title>",
        f"<style>{STYLE}</style></head><body><div class='wrap'>",
        f"<h1>{html.escape(fund)}: annual refresh for review</h1>",
        (
            f"<p class='sub'>Prepared {datetime.now(UTC):%d %B %Y}. Every change below is shown "
            f"against the value it replaces and the cell it came from.</p>"
        ),
        (
            "<div class='tot'>"
            f"<span><b>{len(plan.edits)}</b> changes</span>"
            f"<span><b>{plan.operations}</b> documents touched</span>"
            f"<span><b>{len(plan.unchanged)}</b> fields unchanged</span>"
            f"<span><b>{len(plan.unresolved)}</b> need a decision</span>"
            "</div>"
        ),
    ]

    for document, edits in sorted(plan.by_document().items()):
        parts.append(f"<h2>{html.escape(document)} &middot; {len(edits)} change(s)</h2>")
        for edit in edits:
            parts.append(
                "<div class='card'><div class='row'>"
                f"<span class='label'>{html.escape(edit.label)}</span>"
                f"<span class='doc'>{html.escape(edit.field_id)}</span></div>"
                "<div class='row'>"
                f"<span class='old'>{html.escape(edit.old_value)}</span>"
                "<span class='arrow'>&rarr;</span>"
                f"<span class='new'>{html.escape(edit.new_value)}</span></div>"
                f"<div class='cite'>source: <code>{html.escape(edit.citation)}</code></div></div>"
            )

    if plan.unresolved:
        parts.append("<h2>Needs a human decision</h2>")
        for item in plan.unresolved:
            parts.append(
                "<div class='card'><div class='row'>"
                f"<span class='label'>{html.escape(item.label)}</span>"
                f"<span class='doc'>{html.escape(item.field_id)}</span></div>"
                f"<div class='cite warn'>{html.escape(item.reason)}</div></div>"
            )

    if plan.unchanged:
        parts.append("<h2>Unchanged, and therefore not touched</h2>")
        parts.append(
            "<div class='card'><div class='cite'>"
            + html.escape(", ".join(plan.unchanged))
            + (
                " &mdash; these carried the same value as last year, so no edit was requested "
                "and no operation was spent.</div></div>"
            )
        )

    parts.append("<h2>Layout and wording checks</h2>")
    for document, result in sorted(guards.items()):
        state = "ok" if result.ok else "bad"
        parts.append(
            f"<div class='card'><div class='row'><span class='label'>{html.escape(document)}</span>"
            f"<span class='{state}'>{'passed' if result.ok else 'FAILED'}</span></div><ul>"
        )
        for line in result.failures:
            parts.append(f"<li class='bad'>{html.escape(line)}</li>")
        for line in result.warnings:
            parts.append(f"<li class='warn'>{html.escape(line)}</li>")
        for line in result.passed:
            parts.append(f"<li class='ok'>{html.escape(line)}</li>")
        parts.append("</ul></div>")

    parts.append("</div></body></html>")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(parts), encoding="utf-8")
    return path


def write_changelog_docx(
    plan: Plan, path: Path, decisions: dict[str, str] | None = None, fund: str = "Meridian Global Equity Fund"
) -> Path:
    """The change-log document: every edit beside its source value, so each number ties back."""
    decisions = decisions or {}
    doc = Document()
    doc.add_heading(f"{fund}: annual update change log", level=0)
    doc.add_paragraph(
        f"Generated {datetime.now(UTC):%d %B %Y}. Each row states the value that was replaced, "
        f"the value that replaced it, and the cell in the data file it was read from."
    )

    doc.add_heading("Changes applied", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, head in zip(
        table.rows[0].cells,
        ["Document", "Field", "Previous value", "New value", "Source", "Decision"],
        strict=True,
    ):
        cell.text = head
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for edit in plan.edits:
        cells = table.add_row().cells
        values = [
            edit.document,
            edit.label,
            edit.old_value,
            edit.new_value,
            edit.citation,
            decisions.get(f"{edit.document}:{edit.field_id}", "approved"),
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = str(value)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)

    if plan.unresolved:
        doc.add_heading("Referred for a decision, not applied", level=1)
        for item in plan.unresolved:
            doc.add_paragraph(f"{item.label} ({item.field_id}): {item.reason}", style="List Bullet")

    if plan.unchanged:
        doc.add_heading("Unchanged", level=1)
        doc.add_paragraph(
            "The following carried the same value as the prior year and were deliberately not "
            "edited: " + ", ".join(plan.unchanged) + "."
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
